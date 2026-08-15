"""Report generation service for exporting financial data."""

from io import BytesIO
from datetime import datetime, timedelta
from decimal import Decimal
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db.models import Invoice, Quote, Client, Project, Organization
from app.services.financial_service import FinancialService


class ReportService:
    """Service for generating reports in Excel format."""

    def __init__(self, db: Session):
        self.db = db
        self.financial_service = FinancialService(db)

    def _style_header(self, worksheet, row, columns):
        """Apply header styling to a row."""
        fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        font = Font(color="FFFFFF", bold=True, size=12)
        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_num, column_title in enumerate(columns, 1):
            cell = worksheet.cell(row=row, column=col_num)
            cell.value = column_title
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment

    def _style_title(self, worksheet, row, title):
        """Apply title styling."""
        cell = worksheet.cell(row=row, column=1)
        cell.value = title
        cell.font = Font(bold=True, size=14, color="1e40af")
        cell.alignment = Alignment(horizontal="left", vertical="center")

    def _add_section_spacing(self, worksheet, row):
        """Add spacing between sections."""
        return row + 2

    def generate_financial_report(self, org_id: str) -> BytesIO:
        """Generate comprehensive financial report in Excel."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Reporte Financiero"

        # Set column widths
        worksheet.column_dimensions['A'].width = 25
        worksheet.column_dimensions['B'].width = 20
        worksheet.column_dimensions['C'].width = 20

        row = 1
        # Title
        self._style_title(worksheet, row, "REPORTE FINANCIERO")
        worksheet.merge_cells(f'A{row}:C{row}')
        row += 1

        # Date
        worksheet.cell(row=row, column=1).value = f"Fecha: {datetime.now().strftime('%d/%m/%Y')}"
        row += 2

        # Financial Summary
        self._style_title(worksheet, row, "RESUMEN FINANCIERO")
        row += 1

        summary = self.financial_service.get_financial_summary(org_id)
        summary_data = [
            ["Descripción", "Valor"],
            ["Ingresos Totales", f"RD$ {summary['total_issued']:,.2f}"],
            ["Total Cobrado", f"RD$ {summary['total_paid']:,.2f}"],
            ["Total Pendiente", f"RD$ {summary['total_pending']:,.2f}"],
            ["Total Vencido", f"RD$ {summary['total_overdue']:,.2f}"],
            ["Tasa de Cobranza", f"{summary['collection_rate']:.2f}%"],
            ["Valor Promedio Factura", f"RD$ {summary['average_invoice_value']:,.2f}"],
        ]

        for col_num, (label, value) in enumerate(summary_data, 1):
            cell = worksheet.cell(row=row, column=1)
            cell.value = label
            cell.font = Font(bold=True if row == row else False)

            cell = worksheet.cell(row=row, column=2)
            cell.value = value
            cell.alignment = Alignment(horizontal="right")
            row += 1

        row = self._add_section_spacing(worksheet, row)

        # Top Clients
        self._style_title(worksheet, row, "CLIENTES TOP")
        row += 1

        clients = self.financial_service.get_top_clients_by_revenue(org_id, 10)
        self._style_header(worksheet, row, ["Cliente", "Ingresos", "Facturas"])
        row += 1

        for client in clients:
            worksheet.cell(row=row, column=1).value = client['client_name']
            worksheet.cell(row=row, column=2).value = f"RD$ {client['total_revenue']:,.2f}"
            worksheet.cell(row=row, column=3).value = client['invoice_count']
            worksheet.cell(row=row, column=2).alignment = Alignment(horizontal="right")
            worksheet.cell(row=row, column=3).alignment = Alignment(horizontal="center")
            row += 1

        row = self._add_section_spacing(worksheet, row)

        # Quote Conversion
        self._style_title(worksheet, row, "CONVERSIÓN DE COTIZACIONES")
        row += 1

        conversion = self.financial_service.get_quote_conversion(org_id)
        conversion_data = [
            ["Métrica", "Valor"],
            ["Total Cotizaciones", conversion['total_quotes']],
            ["Enviadas", conversion['sent_quotes']],
            ["Aceptadas", conversion['accepted_quotes']],
            ["Rechazadas", conversion['rejected_quotes']],
            ["Tasa Conversión", f"{conversion['conversion_rate']:.2f}%"],
        ]

        for label, value in conversion_data:
            worksheet.cell(row=row, column=1).value = label
            worksheet.cell(row=row, column=2).value = value
            if isinstance(value, str) and '%' in str(value):
                worksheet.cell(row=row, column=2).alignment = Alignment(horizontal="right")
            row += 1

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_invoices_report(self, org_id: str, start_date: datetime = None, end_date: datetime = None) -> BytesIO:
        """Generate detailed invoices report in Excel."""
        if not start_date:
            start_date = datetime.now() - timedelta(days=90)
        if not end_date:
            end_date = datetime.now()

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Facturas"

        # Set column widths
        widths = [15, 20, 15, 12, 15, 15, 15, 15]
        for idx, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(idx)].width = width

        row = 1
        # Title
        self._style_title(worksheet, row, f"REPORTE DE FACTURAS ({start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')})")
        worksheet.merge_cells(f'A{row}:H{row}')
        row += 2

        # Header
        headers = ["NCF", "Cliente", "Fecha Emisión", "Vencimiento", "Estado", "Subtotal", "Impuesto", "Total"]
        self._style_header(worksheet, row, headers)
        row += 1

        # Data
        invoices = self.db.query(Invoice).filter(
            Invoice.organization_id == org_id,
            Invoice.issue_date >= start_date,
            Invoice.issue_date <= end_date
        ).all()

        total_amount = Decimal(0)
        for invoice in invoices:
            worksheet.cell(row=row, column=1).value = invoice.ncf or "SIN NCF"
            worksheet.cell(row=row, column=2).value = invoice.client.name if invoice.client else "N/A"
            worksheet.cell(row=row, column=3).value = invoice.issue_date.strftime('%d/%m/%Y') if invoice.issue_date else ""
            worksheet.cell(row=row, column=4).value = invoice.due_date.strftime('%d/%m/%Y') if invoice.due_date else ""
            worksheet.cell(row=row, column=5).value = invoice.status
            worksheet.cell(row=row, column=6).value = float(invoice.total - (invoice.tax or Decimal(0))) if invoice.total else 0
            worksheet.cell(row=row, column=7).value = float(invoice.tax or 0)
            worksheet.cell(row=row, column=8).value = float(invoice.total or 0)

            total_amount += invoice.total or Decimal(0)

            # Format numbers
            for col in [6, 7, 8]:
                worksheet.cell(row=row, column=col).number_format = 'RD$ #,##0.00'
            row += 1

        # Totals row
        row += 1
        worksheet.cell(row=row, column=5).value = "TOTAL"
        worksheet.cell(row=row, column=5).font = Font(bold=True)
        worksheet.cell(row=row, column=8).value = float(total_amount)
        worksheet.cell(row=row, column=8).number_format = 'RD$ #,##0.00'
        worksheet.cell(row=row, column=8).font = Font(bold=True)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_quotes_report(self, org_id: str) -> BytesIO:
        """Generate quotes report in Excel."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Cotizaciones"

        # Set column widths
        widths = [15, 15, 12, 15, 15, 15, 12]
        for idx, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(idx)].width = width

        row = 1
        # Title
        self._style_title(worksheet, row, "REPORTE DE COTIZACIONES")
        worksheet.merge_cells(f'A{row}:G{row}')
        row += 2

        # Header
        headers = ["Código", "Cliente", "Estado", "Subtotal", "Descuento", "Impuesto", "Total"]
        self._style_header(worksheet, row, headers)
        row += 1

        # Data
        quotes = self.db.query(Quote).filter(
            Quote.organization_id == org_id
        ).all()

        total_amount = Decimal(0)
        for quote in quotes:
            worksheet.cell(row=row, column=1).value = quote.code
            worksheet.cell(row=row, column=2).value = quote.client.name if quote.client else "N/A"
            worksheet.cell(row=row, column=3).value = quote.status
            worksheet.cell(row=row, column=4).value = float(quote.subtotal or 0)
            worksheet.cell(row=row, column=5).value = float(quote.discount or 0)
            worksheet.cell(row=row, column=6).value = float(quote.tax or 0)
            worksheet.cell(row=row, column=7).value = float(quote.total or 0)

            total_amount += quote.total or Decimal(0)

            # Format numbers
            for col in [4, 5, 6, 7]:
                worksheet.cell(row=row, column=col).number_format = 'RD$ #,##0.00'
            row += 1

        # Totals row
        row += 1
        worksheet.cell(row=row, column=3).value = "TOTAL"
        worksheet.cell(row=row, column=3).font = Font(bold=True)
        worksheet.cell(row=row, column=7).value = float(total_amount)
        worksheet.cell(row=row, column=7).number_format = 'RD$ #,##0.00'
        worksheet.cell(row=row, column=7).font = Font(bold=True)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_clients_report(self, org_id: str) -> BytesIO:
        """Generate clients report with their activity."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Clientes"

        # Set column widths
        widths = [25, 20, 15, 15, 15]
        for idx, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(idx)].width = width

        row = 1
        # Title
        self._style_title(worksheet, row, "REPORTE DE CLIENTES")
        worksheet.merge_cells(f'A{row}:E{row}')
        row += 2

        # Header
        headers = ["Cliente", "Email", "Facturas", "Total Facturado", "Último Movimiento"]
        self._style_header(worksheet, row, headers)
        row += 1

        # Data
        clients = self.db.query(
            Client.id,
            Client.name,
            Client.email,
            func.count(Invoice.id).label('invoice_count'),
            func.sum(Invoice.total).label('total_invoiced')
        ).outerjoin(Invoice, Invoice.client_id == Client.id).filter(
            Client.organization_id == org_id
        ).group_by(Client.id, Client.name, Client.email).all()

        for client in clients:
            worksheet.cell(row=row, column=1).value = client[1]
            worksheet.cell(row=row, column=2).value = client[2] or ""
            worksheet.cell(row=row, column=3).value = client[3] or 0
            worksheet.cell(row=row, column=4).value = float(client[4] or 0)
            worksheet.cell(row=row, column=4).number_format = 'RD$ #,##0.00'
            worksheet.cell(row=row, column=3).alignment = Alignment(horizontal="center")
            row += 1

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
